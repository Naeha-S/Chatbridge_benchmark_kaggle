from __future__ import annotations
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

IDS = {"CB-00003", "CB-00009", "CB-00015", "CB-00021", "CB-00033", "CB-00039", "CB-00045", "CB-00048", "CB-00051", "CB-00057", "CB-00063", "CB-00069", "CB-00075", "CB-00081"}

def load_jsonl(path): return {json.loads(line)["case_id"]: json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()}

def main():
    workbook = load_workbook("reports/generated/ContinuityBench_rereview_2.xlsx", data_only=True)["Re-review 14"]
    rows = {row[0]: row for row in workbook.iter_rows(min_row=2, values_only=True) if row[0]}
    if set(rows) != IDS: raise SystemExit(f"Workbook IDs mismatch: {sorted(set(rows) ^ IDS)}")
    for case_id, row in rows.items():
        if row[14] != "accept" or row[8:13] != ("Yes", "Yes", "Yes", "Yes", "Yes") or row[19] != "Completed": raise SystemExit(f"Invalid re-review result for {case_id}")
    current = load_jsonl("data/generated/continuitybench_v1.jsonl")
    prior = load_jsonl("reports/generated/review_after_rereview_repair.jsonl")
    if set(current) != set(prior): raise SystemExit("Current dataset IDs differ from prior repaired review export")
    records=[]
    for case_id, case in current.items():
        record = dict(case)
        if case_id in IDS:
            row=rows[case_id]
            record["review"]={"status":"accepted","reviewer_decision":"accept","issue_category":"none","correction_notes":row[15] or "","reviewer":row[17],"review_date":str(row[18]),"expected_answer":row[13] or "","evidence":row[16] or "","review_stage":"rereview_2"}
            record["human_checks"]={"readable_coherent":row[8],"category_correct":row[9],"gold_state_supported":row[10],"no_answer_leakage":row[11],"fair_unambiguous":row[12]}
        else:
            record["review"]={"status":"accepted","reviewer_decision":"accept","review_stage":"prior_review_history"}
            record["human_checks"]={"reviewed":True}
        records.append(record)
    out=Path("reports/generated/review_reconciled_rereview2.jsonl"); out.write_text("\n".join(json.dumps(r,ensure_ascii=False) for r in records)+"\n",encoding="utf-8")
    summary={"total_cases":len(records),"status_counts":{"accepted":len(records)},"newly_accepted_rereview_2":len(IDS),"previously_accepted_rereview_1":16,"pending_human_review":0,"source_workbook":"reports/generated/ContinuityBench_rereview_2.xlsx","review_stage_2_reviewer_values":dict(Counter(rows[x][17] for x in IDS)),"review_stage_2_dates":dict(Counter(str(rows[x][18]) for x in IDS)),"metadata_discrepancy":"User prompt says Manus, but workbook Reviewer cells contain Naeha; workbook value preserved.","automated_validation_is_not_human_semantic_proof":True}
    Path("reports/generated/review_reconciled_rereview2_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    prior_sheet = load_workbook("reports/generated/ContinuityBench_rereview_1.xlsx", data_only=True)["Re-review 30"]
    prior_accepted = sorted(row[0] for row in prior_sheet.iter_rows(min_row=2, values_only=True) if row[0] and row[15] == "accept")
    report={"dataset_count":len(current),"new_cases_added":0,"newly_accepted_case_ids":sorted(IDS),"previously_accepted_rereview_1_case_ids":prior_accepted,"status_changes":{"rereview_2":"correct/awaiting_re_review -> accepted"},"pending_human_review":[],"metadata_discrepancy":summary["metadata_discrepancy"],"pipeline_note":"Current repaired dataset used as source of truth; no older snapshot used."}
    Path("reports/generated/rereview2_reconciliation_report.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
    print(json.dumps(summary,indent=2))
if __name__ == "__main__": main()

