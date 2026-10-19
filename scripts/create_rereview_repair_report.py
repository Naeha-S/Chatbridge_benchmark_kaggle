from __future__ import annotations
import json
from pathlib import Path
from openpyxl import load_workbook

def load_jsonl(path): return {json.loads(line)["case_id"]: json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()}

def main():
    workbook = load_workbook("reports/generated/ContinuityBench_rereview_1.xlsx", data_only=True)["Re-review 30"]
    rows = {row[0]: row for row in workbook.iter_rows(min_row=2, values_only=True) if row[0]}
    affected = {case_id: row for case_id, row in rows.items() if row[15] == "correct"}
    accepted = [case_id for case_id, row in rows.items() if row[15] == "accept"]
    current = load_jsonl("data/generated/continuitybench_v1.jsonl")
    history = load_jsonl("reports/generated/review_completed.jsonl")
    outcomes=[]
    for case_id, row in affected.items():
        old = history[case_id]; new = current[case_id]
        outcomes.append({"case_id": case_id, "previous_rereview_decision": "correct", "reviewer": row[18], "review_date": str(row[19]), "repair_status": "awaiting_re_review", "before_task": old["handoff_task"], "after_task": new["handoff_task"], "before_gold_state": old["gold_state"], "after_gold_state": new["expected_continuity_state"], "notes": row[16], "evidence": row[17]})
    report={"source_workbook":"reports/generated/ContinuityBench_rereview_1.xlsx","affected_cases":len(outcomes),"accepted_during_rereview":len(accepted),"accepted_case_ids":accepted,"dataset_count":len(current),"new_cases_added":0,"repair_status":"awaiting_re_review","outcomes":outcomes,"remaining_issues":["The 14 task repairs have not received another human review.","The 16 accepted re-review decisions remain historical and were not modified.","Automated checks do not establish semantic validity."]}
    Path("reports/generated/rereview_task_repair_report.json").write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding="utf-8")
    lines=["# Re-review Task Repair Report","",f"Completed workbook: `reports/generated/ContinuityBench_rereview_1.xlsx`",f"Affected cases: {len(outcomes)}",f"Accepted during re-review and preserved: {len(accepted)}", "Dataset size: 84", "New cases: 0", "Repair status: awaiting another human re-review", ""]
    for item in outcomes: lines += [f"## {item['case_id']}","",f"- Previous decision: `correct`",f"- Reviewer: {item['reviewer']}",f"- Before: {item['before_task']}",f"- After: {item['after_task']}",f"- Repair status: `awaiting_re_review`",f"- Evidence: {item['evidence']}",""]
    lines += ["## Remaining issues","","No affected case was auto-accepted. The 14 changed tasks must be reviewed again by a human."]
    Path("reports/generated/rereview_task_repair_report.md").write_text("\n".join(lines),encoding="utf-8")
    print(json.dumps({"affected_cases":len(outcomes),"accepted_during_rereview":len(accepted),"new_cases_added":0},indent=2))
if __name__ == "__main__": main()
