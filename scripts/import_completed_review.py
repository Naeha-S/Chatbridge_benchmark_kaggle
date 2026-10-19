from __future__ import annotations
import argparse, json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook


def classify(note: str, decision: str) -> str:
    if decision == "accept": return "none"
    text = (note or "").lower()
    if "deadline" in text or "inferred" in text or "provenance" in text: return "unsupported_state"
    return "ambiguous_expected_state"


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--workbook", required=True); parser.add_argument("--review-jsonl", required=True); parser.add_argument("--summary", required=True); parser.add_argument("--output-jsonl", required=True); args = parser.parse_args()
    workbook = load_workbook(args.workbook, data_only=True)
    sheet = workbook["Human Review"]
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        case_id = row[0]
        if not case_id: continue
        rows[case_id] = {"readable": row[9], "category_correct": row[10], "gold_supported": row[11], "no_leakage": row[12], "fair": row[13], "decision": row[14], "expected_answer": row[15], "notes": row[16], "evidence": row[17], "reviewer": row[18], "review_date": row[19], "dataset_status": row[20]}
    records = [json.loads(line) for line in Path(args.review_jsonl).read_text(encoding="utf-8").splitlines() if line.strip()]
    missing = [r["case_id"] for r in records if r["case_id"] not in rows]
    if missing: raise SystemExit(f"Workbook missing case IDs: {missing}")
    for record in records:
        row = rows[record["case_id"]]
        if row["decision"] not in {"accept", "correct"} or row["dataset_status"] != "Reviewed": raise SystemExit(f"Invalid incomplete review for {record['case_id']}")
        record["review"] = {"status": "accepted" if row["decision"] == "accept" else "corrected", "reviewer_decision": row["decision"], "issue_category": classify(row["notes"], row["decision"]), "correction_notes": row["notes"] or "", "reviewer": row["reviewer"], "review_date": str(row["review_date"]), "expected_answer": row["expected_answer"] or "", "evidence": row["evidence"] or ""}
        record["human_checks"] = {"readable_coherent": row["readable"], "category_correct": row["category_correct"], "gold_state_supported": row["gold_supported"], "no_answer_leakage": row["no_leakage"], "fair_unambiguous": row["fair"]}
    output = Path(args.output_jsonl); output.parent.mkdir(parents=True, exist_ok=True); output.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in records) + "\n", encoding="utf-8")
    statuses = Counter(r["review"]["status"] for r in records); issues = Counter(r["review"]["issue_category"] for r in records)
    summary = {"total_cases": len(records), "status_counts": dict(statuses), "issue_category_counts": dict(issues), "invalid_statuses": [], "invalid_issue_categories": [], "semantically_validated": statuses.get("unreviewed", 0) == 0 and statuses.get("corrected", 0) == 0, "source_workbook": str(Path(args.workbook)), "reviewer_counts": dict(Counter(r["review"]["reviewer"] for r in records)), "review_date_counts": dict(Counter(r["review"]["review_date"] for r in records)), "note": "Human decisions imported from the completed workbook; automated checks remain separate."}
    Path(args.summary).write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))

if __name__ == "__main__": main()

